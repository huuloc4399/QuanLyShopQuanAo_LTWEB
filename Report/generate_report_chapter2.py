import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()

    # 1. Page Margins: Left 3cm, Top 3cm, Bottom 3cm, Right 2cm
    for section in doc.sections:
        section.top_margin = Inches(3.0 * 0.393701)
        section.bottom_margin = Inches(3.0 * 0.393701)
        section.left_margin = Inches(3.0 * 0.393701)
        section.right_margin = Inches(2.0 * 0.393701)

    # Base style: Times New Roman 13pt, 1.5 line spacing
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)
    p_format = style.paragraph_format
    p_format.space_before = Pt(0)
    p_format.space_after = Pt(0)
    p_format.line_spacing = 1.5

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(31, 78, 121)
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(13)
            r_bold.font.bold = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.italic = italic
        return p

    def add_table_data(headers, rows):
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Format header
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
            for p in hdr_cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

        # Add rows
        for r_idx, row in enumerate(rows):
            row_cells = table.add_row().cells
            bg_color = "F2F2F2" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = str(val)
                shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color))
                row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)
                for p in row_cells[c_idx].paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    if c_idx in [0, 1, 2]:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(11)

        tblPr = table._tbl.tblPr
        borders = parse_xml(
            r'<w:tblBorders {} >'
            r'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            r'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            r'</w:tblBorders>'.format(nsdecls('w'))
        )
        tblPr.append(borders)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # CONTENT GENERATION
    add_heading_1("CHƯƠNG 2: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")

    add_heading_2("2.1. Sơ đồ Use Case (UML Use Case)")
    add_p("Hệ thống Website Quản lý & Kinh doanh Mặt hàng Quần áo phục vụ 3 nhóm tác nhân chính:")
    add_p("Toàn quyền quản trị hệ thống, quản lý tài khoản và phân quyền chi tiết (RBAC), quản lý danh mục và sản phẩm, duyệt đơn đặt hàng, quản lý kho hàng, thiết lập mã giảm giá (Coupons), xem Dashboard và thống kê doanh thu đa chiều.", bold_prefix="1. Quản trị viên (Admin): ")
    add_p("Thực hiện các nghiệp vụ quản lý kho (kiểm tra lượng tồn theo Size/Màu, lập phiếu nhập hàng từ xưởng may/nhà cung cấp), theo dõi và cập nhật trạng thái đơn hàng (Chờ xác nhận -> Đang giao -> Đã hoàn thành).", bold_prefix="2. Nhân viên (Staff): ")
    add_p("Khách hàng vãng lai hoặc đã đăng ký tài khoản có thể xem sản phẩm, lọc theo giá/danh mục, xem chi tiết kích cỡ/màu sắc/tồn kho, thêm vào giỏ hàng (lưu CSDL và Session), áp mã giảm giá khuyến mãi, đặt hàng trực tuyến, quản lý sổ địa chỉ giao hàng và gửi đánh giá sản phẩm.", bold_prefix="3. Khách hàng (Customer): ")
    add_p("[Hình 2.1: Sơ đồ Use Case tổng quát của Hệ thống Quản lý Quần áo]", italic=True)

    add_heading_2("2.2. Thiết kế Cơ sở dữ liệu (ERD)")
    add_p("Cơ sở dữ liệu WebQuanLyQuanAoDb được thiết kế chuẩn hóa 3NF trên Microsoft SQL Server 2019/2022 gồm 19 bảng, chia thành 5 phân hệ cốt lõi:")
    add_p("Roles, Permissions, RolePermissions, Users, CustomerAddresses.", bold_prefix="1. Phân hệ Người dùng & Phân quyền RBAC: ")
    add_p("Categories, Products, Sizes, Colors, ProductVariants.", bold_prefix="2. Phân hệ Sản phẩm & Biến thể Quần áo: ")
    add_p("Suppliers, ImportReceipts, ImportReceiptDetails.", bold_prefix="3. Phân hệ Quản lý Kho & Nhập hàng: ")
    add_p("Carts, CartItems, Coupons, Orders, OrderDetails.", bold_prefix="4. Phân hệ Giỏ hàng, Khuyến mãi & Đơn hàng: ")
    add_p("ProductReviews.", bold_prefix="5. Phân hệ Tương tác Khách hàng: ")
    add_p("[Hình 2.2: Sơ đồ thực thể liên kết (ERD) Cơ sở dữ liệu WebQuanLyQuanAoDb]", italic=True)

    add_heading_2("2.3. Mô tả chi tiết các bảng dữ liệu")

    # Group 1: RBAC
    add_heading_3("A. Phân hệ Phân quyền & Quản lý Người dùng (RBAC)")

    add_p("Bảng 2.1. Cấu trúc bảng Roles (Vai trò người dùng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["RoleId", "INT IDENTITY(1,1)", "PK", "NO", "Mã vai trò (Khóa chính, tự tăng)"],
            ["RoleName", "NVARCHAR(50)", "UNIQUE", "NO", "Tên vai trò (Admin, Staff, Customer)"],
            ["Description", "NVARCHAR(255)", "", "YES", "Mô tả phạm vi quyền hạn"]
        ]
    )

    add_p("Bảng 2.2. Cấu trúc bảng Permissions (Danh mục quyền chức năng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["PermissionId", "INT IDENTITY(1,1)", "PK", "NO", "Mã quyền (Khóa chính, tự tăng)"],
            ["PermissionName", "NVARCHAR(100)", "", "NO", "Tên quyền hiển thị"],
            ["PermissionCode", "VARCHAR(50)", "UNIQUE", "NO", "Mã code phân quyền (VD: PRODUCT_MANAGE)"],
            ["Module", "NVARCHAR(50)", "", "NO", "Phân hệ trực thuộc (Sản phẩm, Đơn hàng, Kho...)"],
            ["Description", "NVARCHAR(255)", "", "YES", "Mô tả chi tiết quyền"]
        ]
    )

    add_p("Bảng 2.3. Cấu trúc bảng RolePermissions (Gán quyền cho vai trò)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["RolePermissionId", "INT IDENTITY(1,1)", "PK", "NO", "Mã gán quyền (Khóa chính, tự tăng)"],
            ["RoleId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Roles(RoleId)"],
            ["PermissionId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Permissions(PermissionId)"]
        ]
    )

    add_p("Bảng 2.4. Cấu trúc bảng Users (Tài khoản người dùng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["UserId", "INT IDENTITY(1,1)", "PK", "NO", "Mã người dùng (Khóa chính, tự tăng)"],
            ["Username", "VARCHAR(50)", "UNIQUE", "NO", "Tên đăng nhập hệ thống"],
            ["PasswordHash", "VARCHAR(255)", "", "NO", "Mật khẩu đăng nhập"],
            ["FullName", "NVARCHAR(100)", "", "NO", "Họ và tên người dùng"],
            ["Email", "VARCHAR(100)", "", "YES", "Địa chỉ Email"],
            ["PhoneNumber", "VARCHAR(20)", "", "YES", "Số điện thoại liên lạc"],
            ["Address", "NVARCHAR(255)", "", "YES", "Địa chỉ liên hệ"],
            ["Avatar", "VARCHAR(255)", "", "YES", "Đường dẫn ảnh đại diện"],
            ["RoleId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Roles(RoleId)"],
            ["IsActive", "BIT", "", "NO", "Trạng thái (1: Hoạt động, 0: Khóa)"],
            ["CreatedAt", "DATETIME", "", "NO", "Thời điểm tạo tài khoản"]
        ]
    )

    add_p("Bảng 2.5. Cấu trúc bảng CustomerAddresses (Sổ địa chỉ khách hàng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["AddressId", "INT IDENTITY(1,1)", "PK", "NO", "Mã địa chỉ (Khóa chính, tự tăng)"],
            ["UserId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Users(UserId)"],
            ["ReceiverName", "NVARCHAR(100)", "", "NO", "Họ tên người nhận tại địa chỉ này"],
            ["PhoneNumber", "VARCHAR(20)", "", "NO", "Số điện thoại người nhận"],
            ["SpecificAddress", "NVARCHAR(255)", "", "NO", "Địa chỉ chi tiết (số nhà, đường, phường)"],
            ["City", "NVARCHAR(100)", "", "YES", "Tỉnh / Thành phố"],
            ["IsDefault", "BIT", "", "NO", "Địa chỉ mặc định (1: Có, 0: Không)"]
        ]
    )

    # Group 2: Products
    add_heading_3("B. Phân hệ Danh mục, Sản phẩm & Biến thể Quần áo")

    add_p("Bảng 2.6. Cấu trúc bảng Categories (Danh mục sản phẩm)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["CategoryId", "INT IDENTITY(1,1)", "PK", "NO", "Mã danh mục (Khóa chính, tự tăng)"],
            ["CategoryName", "NVARCHAR(100)", "", "NO", "Tên phân loại trang phục"],
            ["Slug", "VARCHAR(100)", "", "YES", "Định danh thân thiện URL"],
            ["Description", "NVARCHAR(500)", "", "YES", "Mô tả nhóm sản phẩm"],
            ["ImageUrl", "VARCHAR(255)", "", "YES", "Hình ảnh đại diện danh mục"],
            ["DisplayOrder", "INT", "", "NO", "Thứ tự sắp xếp trên Menu"],
            ["IsActive", "BIT", "", "NO", "Trạng thái hiển thị (1: Hiện, 0: Ẩn)"]
        ]
    )

    add_p("Bảng 2.7. Cấu trúc bảng Products (Mặt hàng quần áo)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["ProductId", "INT IDENTITY(1,1)", "PK", "NO", "Mã sản phẩm (Khóa chính, tự tăng)"],
            ["CategoryId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Categories(CategoryId)"],
            ["ProductName", "NVARCHAR(200)", "", "NO", "Tên mặt hàng quần áo"],
            ["ProductCode", "VARCHAR(50)", "UNIQUE", "YES", "Mã kiểu dáng (VD: AT-001)"],
            ["Description", "NVARCHAR(MAX)", "", "YES", "Mô tả chất liệu, form dáng, cách bảo quản"],
            ["OriginalPrice", "DECIMAL(18,2)", "", "NO", "Giá vốn / Giá nhập hàng"],
            ["Price", "DECIMAL(18,2)", "", "NO", "Giá bán niêm yết chính thức"],
            ["DiscountPercent", "INT", "", "YES", "% Giảm giá khuyến mãi"],
            ["MainImage", "VARCHAR(255)", "", "YES", "Ảnh đại diện chính của sản phẩm"],
            ["IsFeatured", "BIT", "", "NO", "Sản phẩm nổi bật trang chủ (1: Có)"],
            ["IsActive", "BIT", "", "NO", "Trạng thái kinh doanh (1: Đang bán, 0: Ngừng)"],
            ["CreatedAt", "DATETIME", "", "NO", "Ngày thêm sản phẩm vào hệ thống"]
        ]
    )

    add_p("Bảng 2.8. Cấu trúc bảng Sizes (Kích cỡ quần áo)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["SizeId", "INT IDENTITY(1,1)", "PK", "NO", "Mã kích cỡ (Khóa chính, tự tăng)"],
            ["SizeName", "VARCHAR(20)", "UNIQUE", "NO", "Tên size (S, M, L, XL, XXL, 29, 30...)"],
            ["Description", "NVARCHAR(100)", "", "YES", "Gợi ý cân nặng, chiều cao tương ứng"]
        ]
    )

    add_p("Bảng 2.9. Cấu trúc bảng Colors (Màu sắc sản phẩm)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["ColorId", "INT IDENTITY(1,1)", "PK", "NO", "Mã màu (Khóa chính, tự tăng)"],
            ["ColorName", "NVARCHAR(50)", "UNIQUE", "NO", "Tên màu sắc (Đen, Trắng, Xanh Navy...)"],
            ["ColorHex", "VARCHAR(10)", "", "YES", "Mã màu HEX phục vụ giao diện (#000000)"]
        ]
    )

    add_p("Bảng 2.10. Cấu trúc bảng ProductVariants (Biến thể sản phẩm & Tồn kho)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["VariantId", "INT IDENTITY(1,1)", "PK", "NO", "Mã biến thể (Khóa chính, tự tăng)"],
            ["ProductId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Products(ProductId)"],
            ["SizeId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Sizes(SizeId)"],
            ["ColorId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Colors(ColorId)"],
            ["SKU", "VARCHAR(50)", "UNIQUE", "YES", "Mã định danh quản lý kho (SKU)"],
            ["StockQuantity", "INT", "", "NO", "Số lượng tồn kho thực tế của biến thể"],
            ["VariantImage", "VARCHAR(255)", "", "YES", "Ảnh riêng cho biến thể màu sắc"]
        ]
    )

    # Group 3: Cart, Coupons & Orders
    add_heading_3("C. Phân hệ Giỏ hàng, Mã Giảm Giá & Đơn Đặt Hàng")

    add_p("Bảng 2.11. Cấu trúc bảng Carts (Giỏ hàng người dùng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["CartId", "INT IDENTITY(1,1)", "PK", "NO", "Mã giỏ hàng (Khóa chính, tự tăng)"],
            ["UserId", "INT", "FK", "YES", "Khóa ngoại tham chiếu Users(UserId)"],
            ["SessionId", "VARCHAR(100)", "", "YES", "Mã Session lưu giỏ cho khách vãng lai"],
            ["CreatedAt", "DATETIME", "", "NO", "Thời điểm tạo giỏ hàng"],
            ["UpdatedAt", "DATETIME", "", "NO", "Thời điểm cập nhật giỏ gần nhất"]
        ]
    )

    add_p("Bảng 2.12. Cấu trúc bảng CartItems (Chi tiết giỏ hàng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["CartItemId", "INT IDENTITY(1,1)", "PK", "NO", "Mã chi tiết giỏ (Khóa chính, tự tăng)"],
            ["CartId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Carts(CartId)"],
            ["VariantId", "INT", "FK", "NO", "Khóa ngoại tham chiếu ProductVariants(VariantId)"],
            ["Quantity", "INT", "", "NO", "Số lượng chọn mua (CHECK > 0)"],
            ["AddedAt", "DATETIME", "", "NO", "Thời điểm thêm sản phẩm vào giỏ"]
        ]
    )

    add_p("Bảng 2.13. Cấu trúc bảng Coupons (Mã giảm giá khuyến mãi)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["CouponId", "INT IDENTITY(1,1)", "PK", "NO", "Mã voucher (Khóa chính, tự tăng)"],
            ["CouponCode", "VARCHAR(50)", "UNIQUE", "NO", "Mã khuyến mãi (VD: HUIT2026, FREESHIP)"],
            ["DiscountPercent", "INT", "", "YES", "% Giảm giá trên đơn"],
            ["DiscountAmount", "DECIMAL(18,2)", "", "YES", "Số tiền giảm trực tiếp (VNĐ)"],
            ["MinOrderAmount", "DECIMAL(18,2)", "", "NO", "Giá trị đơn hàng tối thiểu để áp dụng"],
            ["StartDate", "DATETIME", "", "NO", "Ngày bắt đầu hiệu lực"],
            ["EndDate", "DATETIME", "", "NO", "Ngày hết hạn sử dụng"],
            ["UsageLimit", "INT", "", "NO", "Số lượt sử dụng tối đa"],
            ["UsedCount", "INT", "", "NO", "Số lượt đã áp dụng"],
            ["IsActive", "BIT", "", "NO", "Trạng thái hoạt động (1: Kích hoạt, 0: Tắt)"]
        ]
    )

    add_p("Bảng 2.14. Cấu trúc bảng Orders (Đơn đặt hàng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["OrderId", "INT IDENTITY(1,1)", "PK", "NO", "Mã đơn hàng (Khóa chính, tự tăng)"],
            ["UserId", "INT", "FK", "YES", "Khóa ngoại tham chiếu Users(UserId)"],
            ["CouponId", "INT", "FK", "YES", "Khóa ngoại tham chiếu Coupons(CouponId)"],
            ["OrderDate", "DATETIME", "", "NO", "Thời điểm đặt mua"],
            ["ReceiverName", "NVARCHAR(100)", "", "NO", "Họ tên người nhận bưu kiện"],
            ["ReceiverPhone", "VARCHAR(20)", "", "NO", "Số điện thoại nhận hàng"],
            ["ShippingAddress", "NVARCHAR(255)", "", "NO", "Địa chỉ giao hàng tận nơi"],
            ["OrderNotes", "NVARCHAR(500)", "", "YES", "Ghi chú đơn hàng"],
            ["TotalAmount", "DECIMAL(18,2)", "", "NO", "Tổng tiền thanh toán sau giảm giá"],
            ["DiscountAmount", "DECIMAL(18,2)", "", "NO", "Số tiền được giảm từ voucher"],
            ["PaymentMethod", "NVARCHAR(50)", "", "NO", "Hình thức thanh toán (COD, Chuyển khoản...)"],
            ["PaymentStatus", "NVARCHAR(50)", "", "NO", "Trạng thái thanh toán"],
            ["OrderStatus", "NVARCHAR(50)", "", "NO", "Trạng thái đơn (Chờ duyệt, Đang giao, Hoàn thành, Hủy)"]
        ]
    )

    add_p("Bảng 2.15. Cấu trúc bảng OrderDetails (Chi tiết đơn hàng)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["OrderDetailId", "INT IDENTITY(1,1)", "PK", "NO", "Mã chi tiết đơn (Khóa chính, tự tăng)"],
            ["OrderId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Orders(OrderId)"],
            ["VariantId", "INT", "FK", "NO", "Khóa ngoại tham chiếu ProductVariants(VariantId)"],
            ["Quantity", "INT", "", "NO", "Số lượng mặt hàng mua (CHECK > 0)"],
            ["UnitPrice", "DECIMAL(18,2)", "", "NO", "Đơn giá tại thời điểm đặt hàng"],
            ["TotalPrice", "DECIMAL(18,2)", "", "NO", "Thành tiền tự động tính (Quantity * UnitPrice)"]
        ]
    )

    # Group 4: Reviews
    add_heading_3("D. Phân hệ Đánh giá Sản phẩm")

    add_p("Bảng 2.16. Cấu trúc bảng ProductReviews (Đánh giá & Bình luận)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["ReviewId", "INT IDENTITY(1,1)", "PK", "NO", "Mã đánh giá (Khóa chính, tự tăng)"],
            ["ProductId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Products(ProductId)"],
            ["UserId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Users(UserId)"],
            ["Rating", "INT", "", "NO", "Điểm số sao đánh giá (1 đến 5 sao)"],
            ["Comment", "NVARCHAR(1000)", "", "YES", "Nội dung nhận xét của khách hàng"],
            ["CreatedAt", "DATETIME", "", "NO", "Thời điểm đăng đánh giá"],
            ["IsApproved", "BIT", "", "NO", "Trạng thái kiểm duyệt (1: Đã duyệt)"]
        ]
    )

    # Group 5: Suppliers & Imports
    add_heading_3("E. Phân hệ Quản lý Nhà Cung Cấp & Nhập Kho")

    add_p("Bảng 2.17. Cấu trúc bảng Suppliers (Nhà cung cấp / Xưởng may)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["SupplierId", "INT IDENTITY(1,1)", "PK", "NO", "Mã nhà cung cấp (Khóa chính, tự tăng)"],
            ["SupplierName", "NVARCHAR(150)", "", "NO", "Tên đơn vị cung cấp / xưởng may"],
            ["Phone", "VARCHAR(20)", "", "YES", "Số điện thoại liên hệ"],
            ["Email", "VARCHAR(100)", "", "YES", "Email liên hệ"],
            ["Address", "NVARCHAR(255)", "", "YES", "Địa chỉ trụ sở / xưởng may"]
        ]
    )

    add_p("Bảng 2.18. Cấu trúc bảng ImportReceipts (Phiếu nhập kho)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["ImportId", "INT IDENTITY(1,1)", "PK", "NO", "Mã phiếu nhập kho (Khóa chính, tự tăng)"],
            ["SupplierId", "INT", "FK", "NO", "Khóa ngoại tham chiếu Suppliers(SupplierId)"],
            ["UserId", "INT", "FK", "NO", "Nhân viên lập phiếu nhập (Users)"],
            ["ImportDate", "DATETIME", "", "NO", "Thời điểm nhập hàng"],
            ["TotalAmount", "DECIMAL(18,2)", "", "NO", "Tổng giá trị lô hàng nhập"],
            ["Notes", "NVARCHAR(500)", "", "YES", "Ghi chú kiểm kho"]
        ]
    )

    add_p("Bảng 2.19. Cấu trúc bảng ImportReceiptDetails (Chi tiết phiếu nhập kho)")
    add_table_data(
        ["Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Mô tả"],
        [
            ["ImportDetailId", "INT IDENTITY(1,1)", "PK", "NO", "Mã chi tiết phiếu nhập (Khóa chính, tự tăng)"],
            ["ImportId", "INT", "FK", "NO", "Khóa ngoại tham chiếu ImportReceipts(ImportId)"],
            ["VariantId", "INT", "FK", "NO", "Khóa ngoại tham chiếu ProductVariants(VariantId)"],
            ["Quantity", "INT", "", "NO", "Số lượng mặt hàng nhập về kho (CHECK > 0)"],
            ["ImportPrice", "DECIMAL(18,2)", "", "NO", "Đơn giá nhập hàng từ nhà cung cấp"],
            ["TotalPrice", "DECIMAL(18,2)", "", "NO", "Thành tiền tự động tính (Quantity * ImportPrice)"]
        ]
    )

    doc.save("BaoCao_Chuong2_ThietKeCSDL.docx")
    print("Successfully generated updated BaoCao_Chuong2_ThietKeCSDL.docx with 19 tables")

if __name__ == "__main__":
    create_report()
